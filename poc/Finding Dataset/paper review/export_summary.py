"""Export ตารางสรุป paper กลุ่ม A/B จาก paper_review.json เป็น CSV, XLSX และ HTML สำหรับนำเข้า Canva

รัน paper_review.py ก่อน แล้วรัน:  ../../.venv/bin/python export_summary.py
ได้ไฟล์ใน export/:
  summary.csv   UTF-8 มี BOM (เปิดใน Excel/Google Sheets แล้วภาษาไทยไม่เพี้ยน)
  summary.xlsx  คัดลอกช่วงเซลล์ไปวางในตารางของ Canva ได้
  summary.html  หน้าตาราง ใช้ถ่ายภาพ PNG ด้วย browser
"""

import csv
import html
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
OUT = HERE / "export"

HEADER = ["paper", "แหล่งตีพิมพ์", "กลุ่ม", "สถาปัตยกรรม", "ชุดข้อมูล", "ผลเด่น",
          "เทรนเดี่ยว", "รวมศูนย์", "วัดผลการแชร์ได้ไหม"]


def rows(d: dict) -> list[list[str]]:
    papers = {p["key"]: p for p in d["papers"]}
    return [[r["name"], r["venue"], papers[r["key"]]["tier"], r["arch"], ", ".join(papers[r["key"]]["datasets"]),
             r["best"], r["local"], r["central"], r["measurable"]]
            for r in d["summary"]["rows"]]


def write_xlsx(path: Path, body: list[list[str]], counts: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "สรุป"
    ws.append(HEADER)
    for r in body:
        ws.append(r)
    head = PatternFill("solid", fgColor="17212B")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHI", (16, 16, 7, 26, 28, 24, 20, 18, 26)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("ตัวเลข")
    ws2.append(["ประเด็น", "จำนวน"])
    for c in ws2[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for c in counts:
        ws2.append([c["item"], c["value"]])
    ws2.column_dimensions["A"].width, ws2.column_dimensions["B"].width = 44, 44
    wb.save(path)


def write_html(path: Path, body: list[list[str]]) -> None:
    e = html.escape
    tr = "".join(
        "<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>"
        for r in body)
    path.write_text(f"""<!doctype html><meta charset="utf-8">
<link href="https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600&display=swap" rel="stylesheet">
<style>
body {{ margin: 0; padding: 24px; background: #fff; font-family: Sarabun, sans-serif; color: #17212b; }}
h1 {{ font-size: 22px; margin: 0 0 12px; }}
table {{ border-collapse: collapse; width: 1600px; font-size: 14px; }}
th {{ background: #17212b; color: #fff; text-align: left; padding: 8px; font-weight: 600; }}
td {{ border-bottom: 1px solid #dce2e8; padding: 8px; vertical-align: top; }}
tr:nth-child(even) td {{ background: #f4f6f8; }}
</style>
<h1>สรุป paper: swarm learning และงานคล้าย SL บนข้อมูล cyber</h1>
<table><thead><tr>{''.join(f'<th>{e(h)}</th>' for h in HEADER)}</tr></thead><tbody>{tr}</tbody></table>
""", encoding="utf-8")


def main() -> None:
    d = json.loads((HERE / "paper_review.json").read_text(encoding="utf-8"))
    body = rows(d)
    OUT.mkdir(exist_ok=True)
    with open(OUT / "summary.csv", "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows([HEADER, *body])
    write_xlsx(OUT / "summary.xlsx", body, d["summary"]["counts"])
    write_html(OUT / "summary.html", body)
    print(f"wrote {', '.join(p.name for p in sorted(OUT.iterdir()))}")


if __name__ == "__main__":
    main()
